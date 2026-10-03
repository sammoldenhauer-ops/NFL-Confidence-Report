# Usage: python3 add_calculator.py <report.xlsx>   (run after slate_book2.py, before recalc)
# Adds the 'Entry Calculator' tab: beginner walk-through of hit chance, break-even, EV, Power vs Flex.
import sys, pickle, re
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.formatting.rule import CellIsRule
from openpyxl.worksheet.datavalidation import DataValidation

f = sys.argv[1]; wb = load_workbook(f)
if 'Entry Calculator' in wb.sheetnames: del wb['Entry Calculator']
ws = wb.create_sheet('Entry Calculator', wb.sheetnames.index('Combo Builder') + 1)

def F(**k): return Font(name='Arial', size=k.pop('size', 10), **k)
INF = PatternFill('solid', fgColor='FFF2CC'); INFONT = F(color='FF0000FF')
HDR = PatternFill('solid', fgColor='1F3864'); SUB = PatternFill('solid', fgColor='D9E1F2')
PCT = '0.0%'; EV = '+0.0%;-0.0%;0.0%'

def put(cell, v, fmt=None, **fk):
    c = ws[cell]; c.value = v; c.font = F(**fk)
    if fmt: c.number_format = fmt
    return c
def inp(cell, v, fmt=None):
    c = put(cell, v, fmt); c.font = INFONT; c.fill = INF
def header(row, text):
    for col in 'ABCDE': ws[f'{col}{row}'].fill = HDR
    put(f'A{row}', text, bold=True, color='FFFFFFFF', size=11)
def subhdr(row, labels, cols='ABCDE'):
    for col, t in zip(cols, labels):
        c = put(f'{col}{row}', t, bold=True); c.fill = SUB

# ---- example: tonight's top 4-leg from Top Entries ----
E = pd.read_pickle('slate_entries.pkl'); ex = E[E['size'] == 4].sort_values('p_all', ascending=False).iloc[0]
names = [re.sub(r'\s+\[.*$', '', l) for l in ex['legs']]

put('A1', 'Entry Calculator: learn the math', bold=True, size=14)
put('A2', 'Type each leg\'s chance in the yellow cells. Everything below updates, and column B says what the formula is doing in plain words.', italic=True)
put('A3', 'Filled in as an example: tonight\'s top 4-leg from Top Entries. Clear the yellow cells to try your own.', italic=True)

# ---- Step 1: inputs ----
header(5, 'STEP 1: YOUR ENTRY')
subhdr(6, ['Leg', 'Pick (optional note)', 'Chance it hits', 'After safety haircut', 'helper: miss ÷ hit'])
for i in range(6):
    r = 7 + i; put(f'A{r}', f'Leg {i+1}')
    inp(f'B{r}', names[i] if i < len(names) else None)
    inp(f'C{r}', round(ex['leg_probs'][i] / 100, 3) if i < len(names) else None, PCT)
    put(f'D{r}', f'=IF(C{r}="","",IF(C{r}>=0.95,C{r},MAX(C{r}-$C$14,0.01)))', PCT)
    put(f'E{r}', f'=IF(D{r}="",0,(1-D{r})/D{r})', '0.000', color='FF808080')
put('A14', 'Safety haircut per leg'); inp('C14', 0.01, PCT)
put('B14', "The v5 3-point cut was mostly a centering bug fixed in v6; 1 point kept as a buffer until Week 4 grading confirms. Set 0% to see the raw model. Legs typed at 95%+ (promo free spaces) are not cut.", italic=True)
put('A15', 'Stake ($)'); inp('C15', 10, '$#,##0.00')
put('A16', 'All-hit chance from Top Entries (optional)'); inp('C16', round(float(ex['p_all']), 3), PCT)
put('B16', "Copy the 'All hit' % for this exact entry from Top Entries or Combo Builder. It includes correlation. Leave blank if the legs are from different games.", italic=True)
put('A17', 'Payout shown in the app (optional)'); inp('C17', None, '0.00"x"')
put('B17', 'PrizePicks payouts change with the lines you pick and with promos. Type the "To Win" multiplier from the app here; blank = Payouts tab default.', italic=True)
dv = DataValidation(type='decimal', operator='between', formula1='0.01', formula2='0.99', allow_blank=True,
                    showErrorMessage=True, errorTitle='Type a percent', error='Type it as a percent, like 55%')
ws.add_data_validation(dv); dv.add('C7:C12'); dv.add('C16')

# ---- Step 2: the math ----
G = '=IF($C$20<2,"",'  # guard: needs 2+ legs
rows = [
 (20, 'Number of legs', 'Count the legs you filled in.', '=COUNT(D7:D12)', '0'),
 (21, 'Chance ALL legs hit (if independent)', 'Multiply every leg together: leg 1 × leg 2 × leg 3 ...', G + 'PRODUCT(D7:D12))', PCT),
 (22, 'Chance ALL legs hit (with correlation)', 'Your pasted all-hit %, lowered by the same haircut. Blank = same as the row above.', G + 'IF(C16="",C21,C16*C21/PRODUCT(C7:C12)))', PCT),
 (23, 'Power payout', 'Your app payout from row 17 if filled in, otherwise the Payouts tab default.', G + 'IF(C17<>"",C17,INDEX(Payouts!$B$2:$B$6,MATCH(C20,Payouts!$A$2:$A$6,0))))', '0.00"x"'),
 (24, 'Entry break-even', '1 ÷ payout. The whole entry must hit at least this often.', G + '1/C23)', PCT),
 (25, 'Break-even PER LEG', '(Entry break-even) raised to (1 ÷ legs). Each leg needs to hit about this often.', G + 'C24^(1/C20))', PCT),
 (26, 'Your typical leg', '(Chance all hit) raised to (1 ÷ legs). The "average" of your legs, done the multiplying way.', G + 'C21^(1/C20))', PCT),
 (27, 'Cushion per leg', 'Typical leg minus break-even per leg. Positive = an edge, negative = a leak.', G + 'C26-C25)', '+0.0%;-0.0%;0.0%'),
 (28, 'Quick EV estimate (mental math)', 'Legs × cushion ÷ break-even per leg. Close to the exact number below.', G + 'C20*C27/C25)', EV),
 (29, 'Power EV (exact, legs independent)', 'Chance all hit × payout − 1. The average profit per $1 over many entries.', G + 'C21*C23-1)', EV),
 (30, 'Power EV (with correlation)', 'Same formula, using the correlated chance from row 22.', G + 'C22*C23-1)', EV),
 (31, 'Average money back on your stake', 'Stake × (1 + Power EV with correlation). An average over many tries, not one night.', G + 'C15*(1+C30))', '$#,##0.00'),
 (32, 'Chance this Power entry loses', '1 − chance all hit.', G + '1-C22)', PCT),
 (33, 'Chance of losing 10 in a row', '(Chance it loses) raised to 10. This is why each stake stays small.', G + '(1-C22)^10)', PCT),
]
header(18, 'STEP 2: THE MATH, ONE PIECE AT A TIME')
subhdr(19, ['What', 'Formula in plain words', 'Result'], 'ABC')
for r, a, b, fml, fmt in rows:
    put(f'A{r}', a, bold=r in (25, 30)); put(f'B{r}', b); put(f'C{r}', fml, fmt, bold=r in (25, 30))

FX = '=IF($C$20<3,"",'
PAY = lambda col: f'INDEX(Payouts!${col}$2:${col}$6,MATCH(C20,Payouts!$A$2:$A$6,0))'
header(35, 'FLEX: YOU CAN MISS 1 (OR 2 ON 5-6 LEGS) AND STILL GET PAID')
subhdr(36, ['What', 'Formula in plain words', 'Result'], 'ABC')
flex = [
 (37, 'Chance exactly 1 miss', 'Chance all hit × (add up miss÷hit for each leg).', FX + 'C21*SUM(E7:E12))', PCT),
 (38, 'Chance exactly 2 misses', 'Chance all hit × (add up miss÷hit × miss÷hit for every pair of legs).', FX + 'C21*(SUM(E7:E12)^2-SUMSQ(E7:E12))/2)', PCT),
 (39, 'Flex payouts: all / 1 miss / 2 misses', 'From the Payouts tab.', FX + f'TEXT({PAY("C")},"0.##")&"x / "&TEXT({PAY("D")},"0.##")&"x / "&TEXT({PAY("E")},"0.##")&"x")', None),
 (40, 'Flex EV (legs independent)', 'All hit × top payout + 1 miss × its payout + 2 misses × its payout − 1.', FX + f'C21*{PAY("C")}+C37*{PAY("D")}+C38*{PAY("E")}-1)', EV),
 (41, 'Power EV (legs independent), to compare', 'Row 29 again, so you can compare side by side.', FX + 'C29)', EV),
 (42, 'Better play for these legs', 'Same-game legs that move together tilt this further toward Power: they tend to all hit or all miss.', '=IF(C20<3,"Power (Flex needs 3+ legs)",IF(MAX(C30,C40)<0,"Skip: both lose money",IF(C30>=C40,"Power","Flex")))&IF(C17<>""," (Flex uses default payouts - toggle Flex in the app for its actual multipliers)","")', None),
]
for r, a, b, fml, fmt in flex:
    put(f'A{r}', a, bold=r == 42); put(f'B{r}', b); put(f'C{r}', fml, fmt, bold=r == 42)

# ---- Step 3: cheat sheet ----
header(44, 'STEP 3: CHEAT SHEET - EV IF EVERY LEG HITS THE SAME %')
for col in 'FGHIJ': ws[f'{col}44'].fill = HDR
cols = [('B', 2, 'P'), ('C', 3, 'P'), ('D', 3, 'F'), ('E', 4, 'P'), ('F', 4, 'F'), ('G', 5, 'P'), ('H', 5, 'F'), ('I', 6, 'P'), ('J', 6, 'F')]
subhdr(45, ['Each leg hits'] + [f'{n}-leg {"Power" if t=="P" else "Flex"}' for _, n, t in cols], 'ABCDEFGHIJ')
for k, p in enumerate([x / 100 for x in range(50, 63)]):
    r = 46 + k; put(f'A{r}', p, '0%', bold=True)
    for col, n, t in cols:
        m = lambda c: f'INDEX(Payouts!${c}$2:${c}$6,MATCH({n},Payouts!$A$2:$A$6,0))'
        if t == 'P': fml = f'=$A{r}^{n}*{m("B")}-1'
        else: fml = f'=$A{r}^{n}*{m("C")}+{n}*$A{r}^{n-1}*(1-$A{r})*{m("D")}+{n*(n-1)//2}*$A{r}^{n-2}*(1-$A{r})^2*{m("E")}-1'
        put(f'{col}{r}', fml, EV)
put('A59', 'Break-even per leg', bold=True)
for col, n, t in cols:
    if t == 'P': put(f'{col}59', f'=(1/INDEX(Payouts!$B$2:$B$6,MATCH({n},Payouts!$A$2:$A$6,0)))^(1/{n})', PCT, bold=True)
put('A60', 'Green = makes money on average, red = loses. Find your typical leg % on the left and read across.', italic=True)

# ---- formulas to remember ----
header(62, 'FORMULAS TO REMEMBER')
for i, t in enumerate([
 '1. Chance the entry hits = multiply every leg (55% × 55% = 30%).',
 '2. Entry break-even = 1 ÷ payout (4-leg Power: 1 ÷ 10 = 10%).',
 '3. Break-even per leg = look it up in row 59 (2-leg 57.7%, 3-leg 58.5%, 4-leg 56.2%, 5-leg 54.9%, 6-leg 54.7% at default payouts).',
 '4. EV = chance the entry hits × payout − 1. Above 0% = profit on average; below 0% = skip.',
 '5. Quick check: EV ≈ legs × (your leg % − break-even %) ÷ break-even %. Example: 4 legs at 58%: 4 × 1.8 ÷ 56.2 ≈ +13%.',
 '6. Same-game legs that move together hit more often than rule 1 says. Use the All hit % from Top Entries or Combo Builder for those.',
 '7. A positive EV is an average over many entries. Most big entries lose on any single night, so keep each stake small.']):
    put(f'A{63+i}', t)

# ---- formatting ----
green = PatternFill('solid', fgColor='C6EFCE'); red = PatternFill('solid', fgColor='FFC7CE')
for rng in ['C27:C30', 'C40:C41', 'B46:J58']:
    ws.conditional_formatting.add(rng, CellIsRule(operator='greaterThan', formula=['0.0000001'], fill=green))
    ws.conditional_formatting.add(rng, CellIsRule(operator='lessThan', formula=['-0.0000001'], fill=red))
for col, w in zip('ABCDEFGHIJ', [40, 70, 14, 16, 14, 12, 12, 12, 12, 12]): ws.column_dimensions[col].width = w
for row in ws.iter_rows(min_row=20, max_row=42, min_col=2, max_col=2):
    for c in row: c.alignment = Alignment(wrap_text=True, vertical='top')
ws.freeze_panes = 'A5'

# ---- How To Use pointer ----
hu = wb['How To Use']; r = hu.max_row + 2
hu.cell(r, 1, "5. Learn the math: 'Entry Calculator' walks through hit chance, break-even per leg, EV, and Power vs Flex for any legs you type in, plus a cheat sheet.").font = F()
wb.save(f); print('Entry Calculator added; example legs:', names)
