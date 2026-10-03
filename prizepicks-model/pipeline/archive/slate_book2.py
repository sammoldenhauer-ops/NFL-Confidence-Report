import pandas as pd, numpy as np, pickle
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter as CL
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import ColorScaleRule
S=pickle.load(open('slate_sims_centered.pkl','rb')); L=pd.read_pickle('slate_legs_final.pkl'); E=pd.read_pickle('slate_entries.pkl')
RAW=pd.read_pickle('slate_legs.pkl'); G=pd.read_pickle('slate_games.pkl'); I=pickle.load(open('slate_inputs.pkl','rb'))
F=lambda **k:Font(name='Arial',**k); HD=PatternFill('solid',start_color='1F3864'); IN=PatternFill('solid',start_color='FFF2CC'); NS=2000
def hdr(ws,row,vals):
    for j,v in enumerate(vals,1): c=ws.cell(row,j,v); c.font=F(bold=True,color='FFFFFF'); c.fill=HD; c.alignment=Alignment(wrap_text=True,vertical='center')
wb=Workbook(); wb.remove(wb.active)
# ---------- Payouts (editable) ----------
pay=wb.create_sheet('Payouts'); hdr(pay,1,['Legs','Power (all hit)','Flex: all hit','Flex: 1 miss','Flex: 2 misses'])
for i,r in enumerate([[2,3,None,None,None],[3,5,2.25,1.25,0],[4,10,5,1.5,0],[5,20,10,2,0.4],[6,37.5,25,2,0.4]],2):
    for j,v in enumerate(r,1):
        c=pay.cell(i,j,v); c.font=F(color='0000FF' if j>1 else '000000')
        if j>1: c.fill=IN
pay.cell(8,1,"Blue cells = PrizePicks multipliers. Defaults are common values; check your app (they vary by state and promos) and edit here. Every EV in the file recalculates.").font=F(italic=True)
for col,w in zip('ABCDE',[8,16,14,14,15]): pay.column_dimensions[col].width=w
# ---------- Sims ----------
keys=[f"{p} | {s}" for p,s in zip(L['player'],L['stat'])]; keys=list(dict.fromkeys(keys))
sims=wb.create_sheet('Sims'); 
for j,k in enumerate(keys,1): sims.cell(1,j,k)
rows=np.array([S[tuple(k.split(' | '))][:NS] for k in keys]).T.round().astype(int)
for i,row in enumerate(rows,2):
    for j,v in enumerate(row,1): sims.cell(i,j,int(v))
LAST=CL(len(keys))
# ---------- Lists (dropdown source) ----------
li=wb.create_sheet('Lists'); li.cell(1,1,'Player | Stat')
for i,k in enumerate(sorted(keys),2): li.cell(i,1,k)
li.cell(1,2,'Side'); li.cell(2,2,'More'); li.cell(3,2,'Less')
# ---------- Combo Builder ----------
cb=wb.create_sheet('Combo Builder',0)
cb.cell(1,1,'Combo Builder').font=F(bold=True,size=14)
cb.cell(2,1,'Pick up to 6 legs: choose Player | Stat, More or Less, and type the PrizePicks line. Yellow cells are inputs.').font=F(italic=True)
hdr(cb,5,['Leg','Player | Stat','More / Less','PrizePicks line','(col)','Chance this leg hits'])
dv=DataValidation(type='list',formula1=f"=Lists!$A$2:$A${len(keys)+1}",allow_blank=True); dv2=DataValidation(type='list',formula1='"More,Less"',allow_blank=True)
cb.add_data_validation(dv); cb.add_data_validation(dv2)
calc=wb.create_sheet('Calc')
for i in range(6):
    r=6+i; cb.cell(r,1,i+1).font=F()
    for c in (2,3,4): cb.cell(r,c).fill=IN; cb.cell(r,c).font=F(color='0000FF')
    dv.add(cb.cell(r,2)); dv2.add(cb.cell(r,3))
    cb.cell(r,5,f'=IF(B{r}="","",MATCH(B{r},Sims!$1:$1,0))').font=F(color='808080')
    col=CL(i+1)
    cb.cell(r,6,f'=IF(B{r}="","",SUM(Calc!{col}2:{col}{NS+1})/{NS})').number_format='0.0%'
    for s in range(NS):
        rr=s+2
        calc.cell(rr,i+1,f"=IF('Combo Builder'!$B${r}=\"\",1,IF('Combo Builder'!$C${r}=\"Less\",--(INDEX(Sims!$A{rr}:${LAST}{rr},'Combo Builder'!$E${r})<'Combo Builder'!$D${r}),--(INDEX(Sims!$A{rr}:${LAST}{rr},'Combo Builder'!$E${r})>'Combo Builder'!$D${r})))")
for s in range(NS): calc.cell(s+2,7,f"=SUM(A{s+2}:F{s+2})")
for c,t in enumerate(['L1','L2','L3','L4','L5','L6','Sum'],1): calc.cell(1,c,t)
res=[('Legs selected','=COUNTA(B6:B11)','0'),
     ('Chance ALL legs hit (from simulations)',f'=IF(B13<1,"",COUNTIF(Calc!G2:G{NS+1},6)/{NS})','0.0%'),
     ('If the legs were independent','=IF(B13<1,"",PRODUCT(F6:F11))','0.0%'),
     ('Correlation lift (all-hit vs independent)','=IF(B13<2,"",B14/B15-1)','+0%;-0%'),
     ('Power break-even',f'=IF(B13<2,"",1/INDEX(Payouts!B2:B6,MATCH(B13,Payouts!A2:A6,0)))','0.0%'),
     ('Power Play expected profit per $1',f'=IF(B13<2,"",B14*INDEX(Payouts!B2:B6,MATCH(B13,Payouts!A2:A6,0))-1)','+0%;-0%'),
     ('Chance exactly 1 miss',f'=IF(B13<3,"",COUNTIF(Calc!G2:G{NS+1},5)/{NS})','0.0%'),
     ('Chance exactly 2 misses',f'=IF(B13<3,"",COUNTIF(Calc!G2:G{NS+1},4)/{NS})','0.0%'),
     ('Flex Play expected profit per $1','=IF(B13<3,"",B14*INDEX(Payouts!C2:C6,MATCH(B13,Payouts!A2:A6,0))+B19*INDEX(Payouts!D2:D6,MATCH(B13,Payouts!A2:A6,0))+B20*INDEX(Payouts!E2:E6,MATCH(B13,Payouts!A2:A6,0))-1)','+0%;-0%')]
for i,(lab,f,fmt) in enumerate(res,13):
    cb.cell(i,1,lab).font=F(bold=True); c=cb.cell(i,2,f); c.font=F(bold=True); c.number_format=fmt
cb.cell(23,1,"Tip: for PrizePicks' .5 lines, 'More 49.5' means 50+. Legs from different games are simulated independently, same-game legs share the same simulated game.").font=F(italic=True)
for col,w in zip('ABCDEF',[42,40,12,15,6,20]): cb.column_dimensions[col].width=w
# ---------- Top Entries ----------
te=wb.create_sheet('Top Entries',0); te.cell(1,1,'Top Entries - most correlated combinations at lines near the market midpoint').font=F(bold=True,size=13)
te.cell(2,1,'Ranked by chance all legs hit. EV assumes the PrizePicks line equals the sportsbook line shown; use the Combo Builder with the actual PrizePicks line before entering.').font=F(italic=True)
r=4
for size in range(2,7):
    hdr(te,r,['Legs','Leg 1','Leg 2','Leg 3','Leg 4','Leg 5','Leg 6','All hit','If independent','1 miss','2 misses','Power EV','Flex EV']); r+=1
    for e in E[E['size']==size].sort_values('p_all',ascending=False).itertuples():
        te.cell(r,1,size).font=F()
        for j,(l,p) in enumerate(zip(e.legs,e.leg_probs)): te.cell(r,2+j,f"{l}  [{p:.0f}%]").font=F()
        for j,v in enumerate([e.p_all,e.indep,e.p_miss1,e.p_miss2],8): c=te.cell(r,j,float(v)); c.number_format='0.0%'; c.font=F()
        te.cell(r,12,f"=H{r}*INDEX(Payouts!$B$2:$B$6,MATCH(A{r},Payouts!$A$2:$A$6,0))-1").number_format='+0%;-0%'
        if size>=3: te.cell(r,13,f"=H{r}*INDEX(Payouts!$C$2:$C$6,MATCH(A{r},Payouts!$A$2:$A$6,0))+J{r}*INDEX(Payouts!$D$2:$D$6,MATCH(A{r},Payouts!$A$2:$A$6,0))+K{r}*INDEX(Payouts!$E$2:$E$6,MATCH(A{r},Payouts!$A$2:$A$6,0))-1").number_format='+0%;-0%'
        r+=1
    r+=1
te.column_dimensions['A'].width=6
for col in 'BCDEFG': te.column_dimensions[col].width=38
for col in 'HIJKLM': te.column_dimensions[col].width=11
# ---------- Legs ----------
lg=wb.create_sheet('Legs',1)
raw=RAW.set_index(['player','stat'])['model_median']
kick=G.set_index('event')['start'].str[:16].str.replace('T',' ')
cols=['Game','Kickoff (UTC)','Team','Player','Stat','Book line','Market fair line','P(More)','P(Less)','Raw model median','Raw model vs market']
hdr(lg,1,cols)
LL=L.sort_values(['event','team','player','stat'])
for i,x in enumerate(LL.itertuples(),2):
    rm=raw.get((x.player,x.stat),np.nan); gap=(rm-x.fair)
    vals=[x.away+' @ '+x.home,kick.get(x.event,''),x.team,x.player,x.stat,x.line,x.fair,x.p_more,x.p_less,round(float(rm),1),round(float(gap),1)]
    for j,v in enumerate(vals,1):
        c=lg.cell(i,j,v); c.font=F()
        if j in (8,9): c.number_format='0.0%'
lg.conditional_formatting.add(f"H2:I{len(LL)+1}",ColorScaleRule(start_type='num',start_value=0.4,start_color='F8696B',mid_type='num',mid_value=0.5,mid_color='FFFFFF',end_type='num',end_value=0.6,end_color='63BE7B'))
for col,w in zip('ABCDEFGHIJK',[12,17,6,22,13,10,14,9,9,15,17]): lg.column_dimensions[col].width=w
lg.freeze_panes='A2'
# ---------- Likely Out ----------
lo=wb.create_sheet('Likely Out',2); hdr(lo,1,['Team','Player','Recent share %','Type','Treated as'])
for i,x in enumerate(I['flags'].itertuples(),2):
    for j,v in enumerate([x.team,x.player,x.recent_share,x.kind,'OUT (share spread to teammates) - please confirm'],1): lo.cell(i,j,v).font=F()
for col,w in zip('ABCDE',[7,20,14,10,48]): lo.column_dimensions[col].width=w
# ---------- Assumptions ----------
a=wb.create_sheet('Assumptions')
notes=[("Data","Lines: SportsGameOdds consensus of 9 sportsbooks, pulled Sun Sep 27 (morning, before kickoff). Play-by-play through Week 3 (nflverse)."),
("Market centering","Each player's simulated range is centered on the market's fair line (bookmaker margin removed). Tested on 626 graded props: market lines beat our projections for every stat, so the model's job here is the SHAPE of outcomes and how teammates move together, not the center."),
("Near-midpoint rule","Only lines where each side is 40-60% are used, so entries stay near the middle rather than easy, lopsided lines."),
("Correlation","v5 game-script simulation: 5,000 simulated games per matchup. Same-game legs share each simulated game; different games are independent. Tested on 2024 and 2025 QB + top receiver pairs: v5 still slightly understates how often stacks hit together, so the lift shown is, if anything, conservative."),
("Likely out","Players with a recent role but no line this week are treated as out, with their share spread to teammates. Confirm on the Likely Out tab."),
("QB-specific shares","Receiver shares are adjusted where the starting QB differs from recent games, using only games the receiver was active for (this week: Rashid Shaheed with Darnold)."),
("Payouts","Edit on the Payouts tab. Expected profit assumes the PrizePicks line matches the book line shown; enter the actual PrizePicks line in the Combo Builder before playing."),
("Views","By Player: full ladder for every stat. By Stat & Confidence: only lines within 2 steps of the market line (the lines PrizePicks is likely to post). Steps: 10 yds passing, 5 yds rushing/receiving (2 if the line is under 25), 1 for receptions and TDs.")]
hdr(a,1,['Item','Detail'])
for i,(k,v) in enumerate(notes,2):
    a.cell(i,1,k).font=F(bold=True); c=a.cell(i,2,v); c.font=F(); c.alignment=Alignment(wrap_text=True,vertical='top')
a.column_dimensions['A'].width=20; a.column_dimensions['B'].width=110
wb.move_sheet('Payouts',offset=-(len(wb.sheetnames)-4))
calc.sheet_state='hidden'; li.sheet_state='hidden'

from openpyxl.worksheet.table import Table, TableStyleInfo
R=pd.read_pickle('ladder_rows.pkl'); SM=pd.read_pickle('ladder_summary.pkl')
STEP={'Pass Yds':10,'Pass TDs':1,'Receptions':1}
R['Steps from market']=[round((ln-m)/STEP.get(st,5 if m>=25 else 2)) for ln,m,st in zip(R['Line'],R['Market line'],R['Stat'])]
def table_sheet(name,df,pct_cols,widths,pos,tname):
    ws=wb.create_sheet(name,pos); cols=list(df.columns)
    for j,c in enumerate(cols,1): ws.cell(1,j,c)
    for i,row in enumerate(df.itertuples(index=False),2):
        for j,v in enumerate(row,1):
            c=ws.cell(i,j,None if (isinstance(v,float) and np.isnan(v)) else v); c.font=F()
            if cols[j-1] in pct_cols: c.number_format='0.0%'
    t=Table(displayName=tname,ref=f"A1:{CL(len(cols))}{len(df)+1}"); t.tableStyleInfo=TableStyleInfo(name='TableStyleMedium2',showRowStripes=True); ws.add_table(t)
    for j,w in enumerate(widths,1): ws.column_dimensions[CL(j)].width=w
    ws.freeze_panes='A2'
    for pc in pct_cols:
        j=cols.index(pc)+1; ws.conditional_formatting.add(f"{CL(j)}2:{CL(j)}{len(df)+1}",ColorScaleRule(start_type='num',start_value=0.5,start_color='FFFFFF',end_type='num',end_value=0.9,end_color='63BE7B'))
    return ws
bp=R.sort_values(['Player','Stat','Line'])[['Player','Team','Game','Kickoff','Stat','Pick','Line','Confidence','Bucket','Market line','At market line','Steps from market']]
table_sheet('By Player',bp,['Confidence'],[22,6,12,15,13,7,8,12,9,12,14,16],0,'ByPlayer')
bs=R[R['Steps from market'].abs()<=2].sort_values(['Stat','Confidence'],ascending=[True,False])[['Stat','Bucket','Confidence','Player','Team','Game','Kickoff','Pick','Line','Market line','Steps from market']]
table_sheet('By Stat & Confidence',bs,['Confidence'],[13,9,12,22,6,12,15,7,8,12,16],1,'ByStat')
sm=SM[['Player','Team','Game','Kickoff','Stat','Market line','Median outcome','P(More) at market','P(Less) at market','More 60%+ up to','More 70%+ up to','More 80%+ up to','Less 60%+ down to','Less 70%+ down to','Less 80%+ down to']].sort_values(['Player','Stat'])
table_sheet('Player Summary',sm,['P(More) at market','P(Less) at market'],[22,6,12,15,13,12,14,15,15,14,14,14,15,15,15],2,'Summary')
hw=wb.create_sheet('How To Use',0)
steps=["HOW TO USE THIS REPORT",
"",
"1. Look up a player: go to 'By Player', click the filter arrow on Player, and pick a name. You'll see every stat he has, with a ladder of lines.",
"   Each row shows which side is favored at that line (More or Less) and how confident. 'At market line' marks the sportsbook line; PrizePicks is usually at or near it.",
"",
"2. Find the strongest plays in a stat: go to 'By Stat & Confidence', filter Stat (e.g., Rec Yds) and Bucket (e.g., 70-79%).",
"   Only lines within 2 steps of the market line are shown here, so you won't see unrealistically low lines. Filter 'Steps from market' to 0 to see only the market line.",
"",
"3. Quick read on everyone: 'Player Summary' shows one row per player-stat: the market line, the chance of More and Less there,",
"   and the sweet spots, e.g., 'More 70%+ up to 44.5' means he clears 44.5 at least 70% of the time.",
"",
"4. Build entries: 'Top Entries' lists the best 2-6 leg entries (same-game stacks that tend to hit together).",
"   Then check any entry in 'Combo Builder' with the actual PrizePicks lines before playing: it gives the combined chance and expected profit.",
"",
"Reading PrizePicks lines: 'More 49.5' means 50 or more. Confidence = chance that pick hits.",
"Players on the 'Likely Out' tab are treated as out and their targets/carries are given to teammates."]
for i,t in enumerate(steps,1):
    c=hw.cell(i,1,t); c.font=F(bold=(i==1),size=14 if i==1 else 11)
hw.column_dimensions['A'].width=150
wb.save('/mnt/user-data/outputs/Slate_Report.xlsx'); print('saved',len(keys),'sim columns; last col',LAST)
